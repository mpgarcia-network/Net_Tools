import csv
import ipaddress
import unicodedata
from io import BytesIO, StringIO

from openpyxl import load_workbook

from .catalog import resolve_os_driver

# sinonimos de cabecalho -> campo interno
HEADER_MAP = {
    "nome": "name",
    "nome do switch": "name",
    "nome do dispositivo": "name",
    "hostname": "name",
    "sysname": "name",
    "name": "name",
    "device": "name",
    "ip": "ip",
    "ip do switch": "ip",
    "ip de gerencia": "ip",
    "ipv4": "ip",
    "ipv4 address": "ip",
    "address": "ip",
    "endereco": "ip",
    "endereco ip": "ip",
    "vendor": "vendor",
    "fabricante": "vendor",
    "modelo": "model",
    "model": "model",
    "hardware": "model",
    "usuario": "username",
    "user": "username",
    "login": "username",
    "username": "username",
    "senha": "password",
    "password": "password",
    "pass": "password",
    "senha de enable": "enable",
    "senha enable": "enable",
    "enable": "enable",
    "enable password": "enable",
    "secret": "enable",
    "protocolo": "protocol",
    "protocol": "protocol",
    "porta": "port",
    "port": "port",
    "porta ssh": "port",
    "driver": "device_type",
    "device_type": "device_type",
    "tipo": "device_type",
    "os": "os",
    "so": "os",
    "operating system": "os",
    "os version": "os",
    "version": "os",
    "tags": "tags",
    "tag": "tags",
}


def _norm(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = "".join(
        c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c)
    )
    return " ".join(text.split())


def normalize_ip(value: str) -> str:
    """Canonicaliza o IP quando possivel (senao devolve o valor limpo)."""
    v = (value or "").strip()
    if not v:
        return ""
    try:
        return str(ipaddress.ip_address(v))
    except ValueError:
        return v


def _build(header, rows, errors: list[str]) -> list[dict]:
    cols: dict[str, int] = {}
    for idx, cell in enumerate(header):
        key = HEADER_MAP.get(_norm(cell))
        if key and key not in cols:
            cols[key] = idx

    if "name" not in cols or "ip" not in cols:
        errors.append("cabecalhos obrigatorios ausentes: informe ao menos 'Nome' e 'IP'")
        return []

    devices: list[dict] = []
    seen: set[str] = set()
    for rownum, row in enumerate(rows, start=2):
        def get(field: str) -> str:
            idx = cols.get(field)
            if idx is None or idx >= len(row):
                return ""
            value = row[idx]
            if isinstance(value, float) and value.is_integer():
                value = int(value)
            return str(value).strip() if value is not None else ""

        name = get("name")
        ip = normalize_ip(get("ip"))
        if not name and not ip:
            continue
        if not name or not ip:
            errors.append(f"linha {rownum}: nome e IP sao obrigatorios")
            continue
        if ip in seen:
            errors.append(f"linha {rownum}: IP duplicado na planilha ({ip}) ignorado")
            continue
        seen.add(ip)

        protocol = (get("protocol") or "ssh").strip().lower()
        if protocol not in ("ssh", "telnet"):
            protocol = "ssh"
        try:
            port = int(get("port") or (23 if protocol == "telnet" else 22))
        except ValueError:
            port = 23 if protocol == "telnet" else 22
        if not (1 <= port <= 65535):
            port = 23 if protocol == "telnet" else 22

        device_type = get("device_type") or resolve_os_driver(get("os"))

        devices.append(
            {
                "name": name,
                "ip": ip,
                "vendor": get("vendor"),
                "model": get("model"),
                "os": get("os"),
                "device_type": device_type,
                "protocol": protocol,
                "port": port,
                "username": get("username"),
                "password": get("password"),
                "enable": get("enable"),
                "tags": get("tags"),
            }
        )
    return devices


def parse_devices_xlsx(content: bytes) -> tuple[list[dict], list[str]]:
    """Le uma planilha .xlsx e devolve (devices, erros)."""
    errors: list[str] = []
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        return [], [f"nao foi possivel abrir a planilha: {e}"]

    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        wb.close()
        return [], ["planilha vazia"]
    devices = _build(header, rows, errors)
    wb.close()
    return devices, errors


def parse_devices_csv(content: bytes) -> tuple[list[dict], list[str]]:
    """Le um CSV (ex.: export do LibreNMS) e devolve (devices, erros)."""
    errors: list[str] = []
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    reader = csv.reader(StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        return [], ["arquivo vazio"]
    devices = _build(header, reader, errors)
    return devices, errors


def parse_devices(filename: str, content: bytes) -> tuple[list[dict], list[str]]:
    """Detecta o formato (xlsx = ZIP 'PK', senao CSV) e delega."""
    fname = (filename or "").lower()
    if fname.endswith(".csv") or (not fname.endswith(".xlsx") and not content.startswith(b"PK")):
        return parse_devices_csv(content)
    return parse_devices_xlsx(content)
